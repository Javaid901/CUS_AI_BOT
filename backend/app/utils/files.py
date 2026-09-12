"""
backend/app/utils/files.py

Safe file handling for uploads:
  - extension validation
  - size limits
  - filename sanitization (prevents path traversal / clobbering)
  - text extraction dispatch per file type
"""

from __future__ import annotations

import re
import uuid
from pathlib import Path

from app.config import settings
from fastapi import HTTPException

_ALLOWED_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_\-\. ]*$")


def validate_upload(filename: str, size: int) -> str:
    """Validate extension + size. Returns the lowercased extension (no dot)."""
    if not filename:
        raise HTTPException(status_code=400, detail="Missing filename")
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in settings.allowed_extensions_list:
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported file type '.{ext}'. Allowed: {', '.join(settings.allowed_extensions_list)}",
        )
    max_bytes = settings.MAX_UPLOAD_MB * 1024 * 1024
    if size > max_bytes:
        raise HTTPException(
            status_code=413,
            detail=f"File too large ({size} bytes). Max {settings.MAX_UPLOAD_MB} MB.",
        )
    return ext


def sanitize_filename(original: str) -> str:
    """Produce a safe stored filename: <uuid>_<sanitized>.<ext>."""
    base, dot, ext = original.rpartition(".")
    clean = re.sub(r"[^A-Za-z0-9_\- ]", "_", base).strip().replace(" ", "_")
    clean = clean[:60] or "document"
    token = uuid.uuid4().hex[:12]
    ext = ext.lower() if dot else ""
    return f"{token}_{clean}.{ext}" if ext else f"{token}_{clean}"


def ensure_dir(path: str) -> Path:
    p = Path(path)
    p.mkdir(parents=True, exist_ok=True)
    return p


def coerce_uuid(value: str):
    """Safely coerce a string to a UUID, returning the original string if invalid."""
    import uuid
    try:
        return uuid.UUID(value)
    except (ValueError, AttributeError):
        return value


def extract_text(path: str, ext: str) -> list[dict]:
    """
    Extract text from a file. Returns a list of page dicts:
        [{"page": int, "text": str}, ...]
    For non-paginated formats (txt/md/docx may be single-section) page is sequential.
    """
    ext = ext.lower()
    if ext == "pdf":
        return _extract_pdf(path)
    if ext == "docx":
        return _extract_docx(path)
    if ext in ("txt", "md"):
        return _extract_text_plain(path)
    raise HTTPException(status_code=400, detail=f"Unsupported extension: {ext}")


def extract_pages_with_tables(path: str, ext: str) -> list[dict]:
    """Like extract_text but also surfaces DOCX table blocks for the date-sheet
    extraction pipeline. PDF/TXT pages are identical to extract_text; DOCX
    returns the paragraph page plus synthetic table blocks:
        {"page": 1, "type": "table", "rows": [[cell, ...], ...], "section": "table:N"}
    """
    ext = ext.lower()
    if ext == "docx":
        return _extract_docx_with_tables(path)
    return extract_text(path, ext)


def _pdf_word_positions(path: str, page_no: int) -> list[dict] | None:
    """Per-page word bounding boxes (PDF points) from PyMuPDF.

    Used by the date-sheet parser to recover true multi-column table layouts
    the text stream flattens. Returns None (never raises) when PyMuPDF is
    unavailable or the page cannot be read — callers degrade to plain text.
    """
    try:
        import fitz
    except ImportError:  # pragma: no cover - optional dependency
        return None
    try:
        with fitz.open(path) as doc:
            if not (1 <= page_no <= len(doc)):
                return None
            page = doc[page_no - 1]
            words = page.get_text("words")
            return [
                {"x0": w[0], "y0": w[1], "x1": w[2], "y1": w[3], "text": w[4]}
                for w in words
            ]
    except Exception:  # pragma: no cover - malformed pdf
        return None


def _pdf_hlines(path: str, page_no: int) -> list[float] | None:
    """Distinct horizontal grid-line y positions (PDF points) for a page.

    CUS date sheets draw their table borders; these lines give the true row
    boundaries that word y-extents alone cannot. Returns None (never raises)
    when PyMuPDF is unavailable or the page has no drawn lines.
    """
    try:
        import fitz
    except ImportError:  # pragma: no cover - optional dependency
        return None
    try:
        with fitz.open(path) as doc:
            if not (1 <= page_no <= len(doc)):
                return None
            page = doc[page_no - 1]
            ys: list[float] = []
            for d in page.get_drawings():
                for item in d["items"]:
                    if item[0] == "l":
                        p1, p2 = item[1], item[2]
                        if abs(p1.y - p2.y) < 0.5:
                            ys.append((p1.y + p2.y) / 2.0)
                    elif item[0] == "re":
                        r = item[1]
                        ys.append(r.y0)
                        ys.append(r.y1)
            if not ys:
                return None
            ys = sorted(set(round(v, 2) for v in ys))
            merged: list[float] = []
            for v in ys:
                if merged and abs(v - merged[-1]) <= 4.0:
                    continue
                merged.append(v)
            return merged
    except Exception:  # pragma: no cover - malformed pdf
        return None


def _pdf_vlines(path: str, page_no: int) -> list[float] | None:
    """Distinct vertical grid-line x positions (PDF points) for a page.

    Table borders are typically drawn as short per-cell segments, so a column
    boundary only counts when its segments' total coverage is substantial
    (stray short rules from logos etc. are ignored). None when unavailable.
    """
    try:
        import fitz
    except ImportError:  # pragma: no cover - optional dependency
        return None
    try:
        with fitz.open(path) as doc:
            if not (1 <= page_no <= len(doc)):
                return None
            page = doc[page_no - 1]
            coverage: dict[float, float] = {}
            for d in page.get_drawings():
                for item in d["items"]:
                    if item[0] == "l":
                        p1, p2 = item[1], item[2]
                        if abs(p1.x - p2.x) < 0.5:
                            x = round((p1.x + p2.x) / 2.0, 2)
                            coverage[x] = coverage.get(x, 0.0) + abs(p1.y - p2.y)
                    elif item[0] == "re":
                        r = item[1]
                        coverage[round((r.x0 + r.x1) / 2.0, 2)] = coverage.get(
                            round((r.x0 + r.x1) / 2.0, 2), 0.0
                        ) + (r.y1 - r.y0)
            if not coverage:
                return None
            xs = sorted(set(x for x, cov in coverage.items() if cov >= 60.0))
            if len(xs) < 2:
                return None
            merged: list[float] = []
            for v in xs:
                if merged and abs(v - merged[-1]) <= 4.0:
                    merged[-1] = (merged[-1] + v) / 2.0
                    continue
                merged.append(v)
            return merged
    except Exception:  # pragma: no cover - malformed pdf
        return None


def _extract_pdf(path: str) -> list[dict]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise HTTPException(status_code=500, detail="PDF library unavailable") from exc

    pages: list[dict] = []
    try:
        reader = PdfReader(path)
        for i, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            pages.append(
                {
                    "page": i,
                    "text": text,
                    "words": _pdf_word_positions(path, i),
                    "hlines": _pdf_hlines(path, i),
                    "vlines": _pdf_vlines(path, i),
                }
            )
    except Exception as exc:  # pragma: no cover - malformed pdf
        raise HTTPException(status_code=422, detail=f"Could not read PDF: {exc}") from exc
    return pages


def _extract_docx(path: str) -> list[dict]:
    try:
        from docx import Document as DocxDocument
    except ImportError as exc:  # pragma: no cover
        raise HTTPException(status_code=500, detail="DOCX library unavailable") from exc

    try:
        doc = DocxDocument(path)
        paras = [p.text for p in doc.paragraphs if p.text.strip()]
        full = "\n".join(paras)
        # DOCX has no reliable page numbers via python-docx; present as one section.
        return [{"page": 1, "text": full}]
    except Exception as exc:  # pragma: no cover
        raise HTTPException(status_code=422, detail=f"Could not read DOCX: {exc}") from exc


def _extract_docx_with_tables(path: str) -> list[dict]:
    try:
        from docx import Document as DocxDocument
    except ImportError as exc:  # pragma: no cover
        raise HTTPException(status_code=500, detail="DOCX library unavailable") from exc

    try:
        doc = DocxDocument(path)
        paras = [p.text for p in doc.paragraphs if p.text.strip()]
        blocks = [{"page": 1, "text": "\n".join(paras) if paras else ""}]
        for ti, table in enumerate(doc.tables, start=1):
            rows = []
            for row in table.rows:
                cells = [((c.text or "").strip()) for c in row.cells]
                if any(cells):
                    rows.append(cells)
            if rows:
                blocks.append({"page": 1, "type": "table", "rows": rows, "section": f"table:{ti}"})
        return blocks
    except Exception as exc:  # pragma: no cover
        raise HTTPException(status_code=422, detail=f"Could not read DOCX: {exc}") from exc


def _extract_text_plain(path: str) -> list[dict]:
    try:
        with open(path, "r", encoding="utf-8", errors="ignore") as fh:
            text = fh.read()
    except Exception as exc:  # pragma: no cover
        raise HTTPException(status_code=422, detail=f"Could not read file: {exc}") from exc
    return [{"page": 1, "text": text}]
